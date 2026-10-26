import os
from datetime import date
import pandas as pd
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------- Edit these if needed ----------
TITLE = "Product Line Profitability & Margin Performance Analysis"
SUBTITLE = "Nassau Candy Distributor"
AUTHOR = "Gopal Krushna Panda"
ORG = "Centurion University of Technology and Management"
PROGRAM = "Unified Mentor Data Analytics Project"
# ------------------------------------------

DARK = RGBColor(0x1F, 0x3A, 0x4D)
CHARTS = "outputs/charts"

# =====================================================
# LOAD DATA AND COMPUTE EVERYTHING
# =====================================================
prod = pd.read_csv("outputs/product_metrics.csv")
div = pd.read_csv("outputs/division_metrics.csv")
clean = pd.read_csv("outputs/clean_data.csv", parse_dates=["Order Date"])
log = pd.read_csv("outputs/cleaning_log.csv") if os.path.exists("outputs/cleaning_log.csv") else None

def money(x): return f"${x:,.0f}"
def money2(x): return f"${x:,.2f}"

total_sales = div["Sales"].sum()
total_profit = div["Gross Profit"].sum()
total_cost = div["Cost"].sum()
total_units = div["Units"].sum()
margin = total_profit / total_sales * 100
n_prod = len(prod)
d0 = clean["Order Date"].min().strftime("%d %b %Y")
d1 = clean["Order Date"].max().strftime("%d %b %Y")
n_rows = len(clean)

by_profit = prod.sort_values("Gross Profit", ascending=False).reset_index(drop=True)
by_margin = prod.sort_values("Gross Margin %", ascending=False).reset_index(drop=True)
top_p, worst, best_m = by_profit.iloc[0], by_margin.iloc[-1], by_margin.iloc[0]
top_div = div.sort_values("Sales", ascending=False).iloc[0]
weak_div = div.sort_values("Gross Margin %").iloc[0]
small_div = div.sort_values("Sales").iloc[0]

def pareto_n(col):
    t = prod.sort_values(col, ascending=False).reset_index(drop=True)
    cum = t[col].cumsum() / t[col].sum() * 100
    return int((cum < 80).sum() + 1), t

n80_p, tp = pareto_n("Gross Profit")
n80_r, tr = pareto_n("Sales")
top80_names = tp["Product Name"].head(n80_p).tolist()
choc_in_top = int((tp.head(n80_p)["Division"] == "Chocolate").sum())

lever_big = top_div["Sales"] * 0.01
lever_small = worst["Sales"] * 0.10
low_volume = prod[prod["Sales"] < 100]
flags = prod[prod["Action Flag"] != "Healthy"].sort_values("Gross Margin %")

# =====================================================
# WORD HELPERS
# =====================================================
def setup(doc):
    st = doc.styles["Normal"]
    st.font.name = "Calibri"; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.15
    for name, size in (("Heading 1", 16), ("Heading 2", 13)):
        h = doc.styles[name]
        h.font.name = "Calibri"; h.font.size = Pt(size); h.font.bold = True
        h.font.color.rgb = DARK
        rf = h.element.rPr.rFonts
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(a)) is not None:
                del rf.attrib[qn(a)]
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1)
        s.top_margin = s.bottom_margin = Inches(0.9)

def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)

def para(doc, text, bold=False, italic=False, size=None, align=None, color=None, after=None):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    if align: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p

def bullets(doc, items, style="List Bullet"):
    if style == "List Number":
        # manual numbering so each list starts again at 1
        for i, it in enumerate(items, 1):
            p = doc.add_paragraph(f"{i}.  {it}")
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.first_line_indent = Inches(-0.25)
    else:
        for it in items:
            doc.add_paragraph(it, style="List Bullet")

def table(doc, header, rows, font=9):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""
        r = c.paragraphs[0].add_run(str(h)); r.bold = True; r.font.size = Pt(font)
        r.font.color.rgb = RGBColor(255, 255, 255)
        shade(c, "34495E")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            v = str(v)
            run = cells[i].paragraphs[0].add_run(v); run.font.size = Pt(font)
            if i > 0 and v[:1] in "$-0123456789":
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if ri % 2 == 1:
                shade(cells[i], "F2F4F6")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def figure(doc, filename, caption, width=6.0):
    path = os.path.join(CHARTS, filename)
    if not os.path.exists(path):
        para(doc, f"[Chart missing: {filename}]", italic=True)
        return
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(width))
    para(doc, caption, italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x55, 0x55, 0x55))

def page_number(section):
    p = section.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        run._r.append(el)

def toc(doc):
    run = doc.add_paragraph().add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t"); txt.text = "Press Ctrl+A then F9 in Word to build the table of contents."
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, sep, txt, end):
        run._r.append(el)

def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# =====================================================
# RESEARCH PAPER
# =====================================================
doc = Document()
setup(doc)
page_number(doc.sections[0])

# ----- Title page -----
for _ in range(5): doc.add_paragraph()
para(doc, TITLE, bold=True, size=26, align=WD_ALIGN_PARAGRAPH.CENTER, color=DARK, after=4)
para(doc, SUBTITLE, size=18, align=WD_ALIGN_PARAGRAPH.CENTER, after=30)
para(doc, "Research Paper: Exploratory Analysis, Insights and Recommendations",
     italic=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=60)
para(doc, f"Prepared by: {AUTHOR}", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para(doc, ORG, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para(doc, PROGRAM, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para(doc, date.today().strftime("%B %Y"), size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
page_break(doc)

# ----- Contents + Abstract -----
doc.add_heading("Table of Contents", 1)
toc(doc)
page_break(doc)

doc.add_heading("Abstract", 1)
para(doc, (
    f"This study analyses {n_rows:,} cleaned order records of Nassau Candy Distributor, covering "
    f"{d0} to {d1}, to find out which product lines truly drive profit and which weaken the company's margins. "
    f"Total sales of {money(total_sales)} generated a gross profit of {money(total_profit)}, a gross margin of {margin:.1f}%. "
    f"Profit is highly concentrated: {n80_p} of {n_prod} products ({n80_p / n_prod * 100:.0f}%) deliver 80% of profit, "
    f"and the {top_div['Division']} division alone accounts for {top_div['Revenue Share %']:.1f}% of revenue. "
    f"{worst['Product Name']} is the weakest product at a {worst['Gross Margin %']:.1f}% margin, and the "
    f"{weak_div['Division']} division earns {weak_div['Gross Margin %']:.1f}% against {margin:.1f}% for the company overall. "
    "The paper closes with pricing, sourcing and portfolio recommendations."))

# ----- 1 Introduction -----
doc.add_heading("1. Introduction", 1)
doc.add_heading("1.1 Background", 2)
para(doc, ("For a distributor, sales volume alone can be misleading. Some products sell in large quantities but earn "
           "little profit, consume a disproportionate share of manufacturing cost, or appear successful while quietly "
           "weakening the overall margin. To stay competitive, Nassau Candy needs to know which products truly drive "
           "profit, which divisions underperform, and where pricing, sourcing or portfolio decisions are needed."))
doc.add_heading("1.2 Problem statement", 2)
para(doc, ("The organisation lacks clear visibility into which product lines deliver the highest gross margin, whether "
           "high-sales products are actually profitable, how profitability varies across divisions, and which products "
           "represent margin risk. Without this, decisions on pricing, promotions and the product portfolio remain "
           "reactive and based on intuition."))
doc.add_heading("1.3 Objectives", 2)
bullets(doc, [
    "Measure gross margin, profit per unit and profit contribution for every product.",
    "Rank products and identify high-profit, high-sales/low-margin and low-sales/low-profit groups.",
    "Compare divisions on revenue, profit and margin, and detect structural margin problems.",
    "Measure profit concentration using Pareto (80/20) analysis and identify over-dependency risk.",
    "Diagnose cost structure and flag products for repricing, cost renegotiation or discontinuation review.",
    "Deliver an interactive Streamlit dashboard so stakeholders can explore the results themselves.",
])

# ----- 2 Data and Method -----
doc.add_heading("2. Data and Methodology", 1)
doc.add_heading("2.1 Dataset", 2)
para(doc, (f"The dataset contains order-level records with order and shipping dates, customer geography, division, "
           f"product, sales, units, cost and gross profit. After cleaning, {n_rows:,} records remain, covering "
           f"{n_prod} products in {len(div)} divisions ({', '.join(div['Division'])}) between {d0} and {d1}."))
doc.add_heading("2.2 Data cleaning and validation", 2)
para(doc, "The following steps were applied before any analysis, using Python (pandas):")
bullets(doc, [
    "Column names and text labels (division, product, region) were trimmed and standardised, including consistent product-name formatting.",
    "Exact duplicate records were removed.",
    "Records with invalid dates, or with missing sales, cost or gross profit, were removed.",
    "Records with zero or negative sales, or negative cost, were removed.",
    "Missing or invalid unit values were filled with the median units for the same product.",
    "Gross profit was cross-checked against Sales minus Cost.",
])
if log is not None:
    para(doc, "Table 1 summarises the effect of each cleaning step.", after=2)
    table(doc, ["Cleaning step", "Rows"], [[r["Step"], f"{int(r['Rows']):,}"] for _, r in log.iterrows()])
    para(doc, "Table 1: Data cleaning summary", italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_heading("2.3 Metrics and KPIs", 2)
table(doc, ["KPI", "Definition"], [
    ["Gross Margin (%)", "Gross Profit ÷ Sales × 100"],
    ["Profit per Unit", "Gross Profit ÷ Units"],
    ["Revenue Contribution", "Product sales ÷ total sales"],
    ["Profit Contribution", "Product gross profit ÷ total gross profit"],
    ["Margin Volatility", "Standard deviation of the product's monthly gross margin"],
])
para(doc, ("Products are grouped by comparing them with the portfolio: 'High profit / High margin' products have profit at "
           "or above the median and margin at or above the overall margin; 'High sales / Low margin' products have sales at "
           "or above the median but margin below the overall margin; 'Low sales / Low profit' products fall below the median "
           "on both."))

# ----- 3 Overall -----
doc.add_heading("3. Overall Performance", 1)
table(doc, ["Metric", "Value"], [
    ["Total sales", money2(total_sales)],
    ["Total cost", money2(total_cost)],
    ["Total gross profit", money2(total_profit)],
    ["Overall gross margin", f"{margin:.2f}%"],
    ["Units sold", f"{total_units:,.0f}"],
    ["Profit per unit", money2(total_profit / total_units)],
    ["Products analysed", str(n_prod)],
])
para(doc, "Table 2: Company-level key performance indicators", italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

# ----- 4 Product -----
doc.add_heading("4. Product-Level Profitability", 1)
doc.add_heading("4.1 Highest-profit products", 2)
para(doc, (f"{top_p['Product Name']} is the largest profit contributor with {money(top_p['Gross Profit'])} "
           f"({top_p['Profit Contribution %']:.1f}% of total profit). Table 3 lists the five most profitable products."), after=2)
table(doc, ["Product", "Division", "Sales", "Gross Profit", "Margin %", "Profit/Unit"],
      [[r["Product Name"], r["Division"], money(r["Sales"]), money(r["Gross Profit"]),
        f"{r['Gross Margin %']:.1f}", money2(r["Profit per Unit"])] for _, r in by_profit.head(5).iterrows()])
para(doc, "Table 3: Top five products by gross profit", italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
figure(doc, "01_profit_by_product.png", "Figure 1: Gross profit by product")

doc.add_heading("4.2 Margin leaderboard", 2)
para(doc, (f"{best_m['Product Name']} has the highest margin at {best_m['Gross Margin %']:.1f}%, while "
           f"{worst['Product Name']} has the lowest at {worst['Gross Margin %']:.1f}%. Table 4 shows the five weakest "
           f"margins; Figure 2 marks in red every product below the overall margin of {margin:.1f}%."), after=2)
table(doc, ["Product", "Division", "Sales", "Gross Profit", "Margin %", "Cost % of Sales"],
      [[r["Product Name"], r["Division"], money(r["Sales"]), money(r["Gross Profit"]),
        f"{r['Gross Margin %']:.1f}", f"{r['Cost % of Sales']:.1f}"]
       for _, r in by_margin.tail(5).iloc[::-1].iterrows()])
para(doc, "Table 4: Five lowest-margin products", italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
figure(doc, "02_margin_by_product.png", "Figure 2: Gross margin % by product (red = below overall margin)")

doc.add_heading("4.3 Product groups", 2)
cat_rows = [[c, "; ".join(g["Product Name"])] for c, g in prod.groupby("Category")]
table(doc, ["Group", "Products"], cat_rows)
para(doc, "Table 5: Product classification", italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

# ----- 5 Division -----
doc.add_heading("5. Division-Level Performance", 1)
table(doc, ["Division", "Sales", "Gross Profit", "Margin %", "Revenue Share %", "Profit Share %"],
      [[r["Division"], money(r["Sales"]), money(r["Gross Profit"]), f"{r['Gross Margin %']:.1f}",
        f"{r['Revenue Share %']:.1f}", f"{r['Profit Share %']:.1f}"] for _, r in div.iterrows()])
para(doc, "Table 6: Division performance", italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
para(doc, (f"The {top_div['Division']} division generates {top_div['Revenue Share %']:.1f}% of revenue and "
           f"{top_div['Profit Share %']:.1f}% of profit, so it carries the business. The {weak_div['Division']} division "
           f"has the weakest margin ({weak_div['Gross Margin %']:.1f}%) and its profit share "
           f"({weak_div['Profit Share %']:.1f}%) is below its revenue share ({weak_div['Revenue Share %']:.1f}%), which "
           f"points to a structural margin issue. The {small_div['Division']} division is very small, with only "
           f"{money(small_div['Sales'])} in sales ({small_div['Revenue Share %']:.2f}% of revenue)."))
figure(doc, "03_division_revenue_vs_profit.png", "Figure 3: Revenue versus gross profit by division", 5.0)
figure(doc, "04_margin_distribution_division.png", "Figure 4: Distribution of order-level margin by division", 5.0)

# ----- 6 Pareto -----
doc.add_heading("6. Profit Concentration (Pareto) Analysis", 1)
para(doc, (f"{n80_r} of {n_prod} products ({n80_r / n_prod * 100:.0f}%) generate 80% of revenue, and {n80_p} of {n_prod} "
           f"({n80_p / n_prod * 100:.0f}%) generate 80% of profit. The products behind 80% of profit are: "
           f"{', '.join(top80_names)}. {choc_in_top} of these {n80_p} products belong to the Chocolate division. "
           f"The single largest product contributes {tp['Gross Profit'].iloc[0] / tp['Gross Profit'].sum() * 100:.1f}% "
           f"of total profit. This is a clear over-dependency risk: a supply, quality or demand problem affecting a few "
           f"chocolate products would hit the whole business."))
figure(doc, "06_pareto_revenue.png", "Figure 5: Pareto chart of revenue")
figure(doc, "07_pareto_profit.png", "Figure 6: Pareto chart of gross profit")

# ----- 7 Cost -----
doc.add_heading("7. Cost Structure Diagnostics", 1)
para(doc, ("Figure 7 plots total cost against total sales for each product. Points far above the dashed break-even line "
           "have cost close to or above sales; points well below it are healthy. Bubble size shows units sold."))
figure(doc, "05_cost_vs_sales.png", "Figure 7: Cost versus sales by product")
para(doc, ("Products with a margin below the overall margin are flagged: 'Cost renegotiation' when sales are at or above the "
           "median (the product sells well but costs too much), 'Discontinuation review' when both sales and profit are "
           "below the median, and 'Repricing' otherwise."), after=2)
table(doc, ["Product", "Division", "Sales", "Margin %", "Flag"],
      [[r["Product Name"], r["Division"], money(r["Sales"]), f"{r['Gross Margin %']:.1f}", r["Action Flag"]]
       for _, r in flags.iterrows()])
para(doc, "Table 7: Products flagged for action", italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

# ----- 8 Insights -----
doc.add_heading("8. Key Insights", 1)
bullets(doc, [
    f"Profit comes from a few products: {n80_p} products ({n80_p / n_prod * 100:.0f}% of the range) deliver 80% of profit, and {choc_in_top} of them are Chocolate products.",
    f"High overall margin hides weak spots: the company margin is {margin:.1f}%, but {worst['Product Name']} earns only {worst['Gross Margin %']:.1f}% (cost is {worst['Cost % of Sales']:.1f}% of its sales).",
    f"The {weak_div['Division']} division earns {weak_div['Gross Margin %']:.1f}% margin versus {top_div['Gross Margin %']:.1f}% in {top_div['Division']}, so mix shift toward {weak_div['Division']} would dilute total margin.",
    f"Size of the prize: a 1-point margin gain in {top_div['Division']} is worth about {money(lever_big)}, more than a 10-point gain on {worst['Product Name']} (about {money(lever_small)}). The biggest profit lever is the core range, not the weakest product.",
    f"{len(low_volume)} products each sold under $100 in total. Their low profit comes mainly from low volume, so the question for them is demand, not just pricing.",
], style="List Number")

# ----- 9 Recommendations -----
doc.add_heading("9. Recommendations", 1)
bullets(doc, [
    f"Protect the core: secure supply, quality and stock levels for {', '.join(top80_names)}, since they generate 80% of profit.",
    f"Reduce concentration risk: build second-tier growth by promoting mid-performing products and testing new items, so profit does not rest on {n80_p} products.",
    f"Fix {worst['Product Name']}: review its selling price and renegotiate its manufacturing cost, then decide whether to keep it if its margin stays near {worst['Gross Margin %']:.0f}%.",
    f"Review the {weak_div['Division']} division: compare prices and unit costs product by product, and target a margin closer to the company average.",
    f"Decide on low-volume products ({', '.join(low_volume['Product Name']) if len(low_volume) else 'none'}): either invest in demand (promotion, wider distribution) or rationalise them, rather than leaving them unmanaged.",
    "Monitor margin monthly using the dashboard and its margin-threshold slider, so that products drifting below the target are caught early.",
], style="List Number")

# ----- 10 Limitations -----
doc.add_heading("10. Limitations", 1)
bullets(doc, [
    "Only gross profit (Sales minus manufacturing cost) is available. Shipping, marketing and overhead costs are not included, so net profit may differ.",
    "Margins are high compared with typical confectionery benchmarks; it is worth confirming how 'Cost' is defined in the source system.",
    "Flags and product groups are relative to the portfolio average and median, so products only slightly below average (for example, within about one percentage point) may be flagged. The dashboard threshold slider lets users apply their own target.",
    "Some products have very small sales, so their margins and volatility figures are based on little data and should be treated with caution.",
    "The analysis covers one dataset and period; seasonality and trends beyond this window are not assessed.",
])

# ----- 11 Conclusion -----
doc.add_heading("11. Conclusion", 1)
para(doc, (f"Nassau Candy is profitable overall ({margin:.1f}% gross margin), but its profit is concentrated in a handful of "
           f"chocolate products, and a few products and one division earn clearly lower margins. Acting on the actions above "
           f"(protecting the core range, fixing the weakest product and division, and managing low-volume items) would move "
           f"the company from reactive to data-driven decisions. The accompanying Streamlit dashboard allows these "
           f"findings to be monitored and updated."))

# ----- Appendices -----
page_break(doc)
doc.add_heading("Appendix A: Full product metrics", 1)
table(doc, ["Product", "Sales", "Units", "Gross Profit", "Margin %", "Profit/Unit", "Volatility"],
      [[r["Product Name"], money(r["Sales"]), f"{r['Units']:,.0f}", money(r["Gross Profit"]),
        f"{r['Gross Margin %']:.1f}", money2(r["Profit per Unit"]), f"{r['Margin Volatility']:.1f}"]
       for _, r in by_profit.iterrows()], font=8)
doc.add_heading("Appendix B: Streamlit dashboard", 1)
para(doc, "The interactive dashboard (app.py) provides these modules:")
bullets(doc, [
    "Insights: automatic findings, margin gauge and a sales-and-margin treemap.",
    "Products: margin leaderboard, profit contribution, profit per unit and product table.",
    "Divisions: revenue versus profit, revenue and profit share, and margin distribution.",
    "Cost vs Margin: cost-sales scatter plot, flag summary and margin-risk list.",
    "Concentration: Pareto charts and dependency indicators.",
    "Trends: monthly trend, division-by-month margin heatmap, region and state views.",
    "Data: downloadable filtered results.",
])
para(doc, "Users can filter by date range, division and region, set a margin threshold, and search for products. Run it with: streamlit run app.py")

doc.save("outputs/Nassau_Profitability_Research_Paper.docx")

# =====================================================
# EXECUTIVE SUMMARY (one page)
# =====================================================
ex = Document()
setup(ex)
ex.styles["Normal"].font.size = Pt(10.5)
ex.styles["Normal"].paragraph_format.space_after = Pt(4)
for s in ex.sections:
    s.top_margin = s.bottom_margin = Inches(0.7)

para(ex, "Executive Summary", bold=True, size=22, color=DARK, after=0)
para(ex, f"{SUBTITLE}: Product Line Profitability & Margin Performance",
     size=12, after=2)
para(ex, f"{AUTHOR} | {date.today().strftime('%B %Y')}", italic=True, size=9, after=8,
     color=RGBColor(0x55, 0x55, 0x55))

ex.add_heading("Purpose", 2)
para(ex, ("This review shows which products and divisions really drive profit at Nassau Candy, and where pricing, "
          "sourcing or portfolio decisions are needed. It is based on "
          f"{n_rows:,} cleaned orders from {d0} to {d1}."))

ex.add_heading("Headline numbers", 2)
table(ex, ["Total sales", "Gross profit", "Gross margin", "Units sold"],
      [[money(total_sales), money(total_profit), f"{margin:.1f}%", f"{total_units:,.0f}"]], font=10)

ex.add_heading("What we found", 2)
bullets(ex, [
    f"Profit is concentrated: {n80_p} of {n_prod} products deliver 80% of profit, and {choc_in_top} of them are Chocolate. The {top_div['Division']} division provides {top_div['Revenue Share %']:.0f}% of revenue.",
    f"{worst['Product Name']} has the weakest margin ({worst['Gross Margin %']:.1f}%); manufacturing cost takes {worst['Cost % of Sales']:.0f}% of its sales.",
    f"The {weak_div['Division']} division earns {weak_div['Gross Margin %']:.1f}% margin versus {margin:.1f}% company-wide.",
    f"{len(low_volume)} products sold under $100 each; their low profit is mainly a volume issue.",
])

ex.add_heading("Recommended actions", 2)
bullets(ex, [
    "Protect supply and quality of the top products that generate most of the profit.",
    f"Reprice or renegotiate cost for {worst['Product Name']}, and review the {weak_div['Division']} division product by product.",
    "Reduce dependency on a few products by growing mid-performing items.",
    "Decide whether to promote or rationalise low-volume products.",
    "Track margins monthly in the live dashboard, using a margin threshold to catch problems early.",
], style="List Number")

para(ex, "Details, charts and limitations are in the full research paper. Figures cover gross profit only (Sales minus manufacturing cost).",
     italic=True, size=9, color=RGBColor(0x55, 0x55, 0x55))
ex.save("outputs/Executive_Summary.docx")

print("Done! Created:")
print(" - outputs/Nassau_Profitability_Research_Paper.docx")
print(" - outputs/Executive_Summary.docx")